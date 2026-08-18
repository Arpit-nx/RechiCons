#Excel manager
from pathlib import Path

from openpyxl import (
    load_workbook,
    Workbook,
)

BASE_DIR = Path(__file__).resolve().parents[2]

class WorkbookManager:

    FILE_PATH = (
        BASE_DIR
        / "storage"
        / "RechiCMS.xlsx"
    )

    def __init__(self):

        self.workbook = None

    def load(self):

        if not self.FILE_PATH.exists():

            self.create()

        self.workbook = load_workbook(
            self.FILE_PATH
        )

        return self.workbook

    def create(self):

        # Create storage directory if it doesn't exist
        self.FILE_PATH.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        from app.repository.schema import SCHEMA #Import for the data schema structure.

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for sheet_name, columns in SCHEMA.items():

            sheet = workbook.create_sheet(sheet_name)

            sheet.append(columns)

        workbook.save(self.FILE_PATH)

    def save(self):

        self.workbook.save(
            self.FILE_PATH
        )
        print("Saving to:", self.FILE_PATH.resolve())
        print("Workbook Sheets:", self.workbook.sheetnames)

    def sheet(self, sheet_name: str):

        if self.workbook is None:
            self.load()

        if sheet_name not in self.workbook.sheetnames:
            self.workbook.create_sheet(sheet_name)
            self.save()

        return self.workbook[sheet_name]

    #Function for gettig the id as identity type.
    def next_id(
        self,
        sheet_name: str,
    ):

        sheet = self.sheet(sheet_name)

        ids = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if row[0] is not None:

                ids.append(int(row[0]))

        return max(ids, default=0) + 1

    #Closing Function.
    def close(self):

        if self.workbook:

            self.workbook.close()

    #Validation function for checking the schema and creating the missing sheets.
    def validate(self):

        from app.repository.schema import SCHEMA

        self.load()

        for sheet in SCHEMA:

            if sheet not in self.workbook.sheetnames:

                ws = self.workbook.create_sheet(sheet)

                ws.append(SCHEMA[sheet])

        self.save()

    #Reading all rows.
    def read_all(
        self,
        sheet_name: str,
    ):

        sheet = self.sheet(sheet_name)

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        data = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            data.append(
                dict(zip(headers, row))
            )

        return data